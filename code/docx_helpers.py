from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
NAVY=RGBColor(0x1f,0x4e,0x79)
def new_doc():
    d=Document(); sec=d.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7)
    for a in ('left_margin','right_margin'): setattr(sec,a,Cm(2.2))
    sec.top_margin=Cm(2.0); sec.bottom_margin=Cm(2.0)
    st=d.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10.5); st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=1.12
    for n,sz in (('Heading 1',17),('Heading 2',13),('Heading 3',11)):
        h=d.styles[n]; h.font.name='Calibri'; h.font.size=Pt(sz); h.font.bold=True; h.font.color.rgb=NAVY
        h.paragraph_format.space_before=Pt(14 if n=='Heading 1' else 10); h.paragraph_format.space_after=Pt(4)
    # footer page number
    p=sec.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run()
    for t,txt in (('begin',None),(None,'PAGE'),('end',None)):
        if t: e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),t)
        else: e=OxmlElement('w:instrText'); e.set(qn('xml:space'),'preserve'); e.text=txt
        r._r.append(e)
    return d
def H(d,t,l=1): return d.add_heading(t,level=l)
def P(d,t,bold=False,italic=False,size=None,align=None):
    p=d.add_paragraph(); _runs(p,t,bold,italic,size)
    if align=='c': p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return p
def _runs(p,t,bold=False,italic=False,size=None):
    parts=t.split('**')
    for i,s in enumerate(parts):
        if not s: continue
        r=p.add_run(s); r.bold=bold or (i%2==1); r.italic=italic
        if size: r.font.size=Pt(size)
def B(d,items,style='List Bullet'):
    for t in items:
        p=d.add_paragraph(style=style); _runs(p,t); p.paragraph_format.space_after=Pt(2)
def shade(cell,hexcol):
    tcPr=cell._tc.get_or_add_tcPr(); s=OxmlElement('w:shd'); s.set(qn('w:val'),'clear'); s.set(qn('w:color'),'auto'); s.set(qn('w:fill'),hexcol); tcPr.append(s)
def T(d,header,rows,widths=None,size=8.5,note=None):
    t=d.add_table(rows=1,cols=len(header)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(header):
        c=t.rows[0].cells[i]; c.text=''; r=c.paragraphs[0].add_run(str(h)); r.bold=True; r.font.size=Pt(size); r.font.color.rgb=RGBColor(255,255,255); shade(c,'1F4E79')
    for k,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=''; _runs(cells[i].paragraphs[0],str(v),False,False,size)
            if i>0 and isinstance(v,str) and v.replace('*','').replace(',','').replace('.','').replace('-','').replace('+','').replace('%','').replace('−','').isdigit(): cells[i].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
            if k%2==1: shade(cells[i],'EEF3F8')
    trPr=t.rows[0]._tr.get_or_add_trPr(); th=OxmlElement('w:tblHeader'); th.set(qn('w:val'),'true'); trPr.append(th)
    for row in t.rows:
        cs=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cs)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Cm(w)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs: p.paragraph_format.space_after=Pt(1)
    if note: p=d.add_paragraph(); r=p.add_run(note); r.italic=True; r.font.size=Pt(8)
    else: d.add_paragraph().paragraph_format.space_after=Pt(2)
    return t
def F(d,path,cap,w=16):
    d.add_picture(path,width=Cm(w)); d.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=d.add_paragraph(); r=p.add_run(cap); r.italic=True; r.font.size=Pt(8.5); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
def code(d,txt,size=7.5):
    for line in txt.rstrip('\n').split('\n'):
        p=d.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0; p.paragraph_format.left_indent=Cm(0.3)
        r=p.add_run(line if line else ' '); r.font.name='Consolas'; r.font.size=Pt(size); r._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Consolas')
def callout(d,t,fill='FFF4E5'):
    tb=d.add_table(rows=1,cols=1); tb.style='Table Grid'; c=tb.rows[0].cells[0]; c.text=''; shade(c,fill); _runs(c.paragraphs[0],t,size=9.5); d.add_paragraph().paragraph_format.space_after=Pt(2)
